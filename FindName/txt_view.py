# txt_view.py
import os
from PySide6.QtGui import QPixmap
from PySide6.QtCore import Qt
import openpyxl

def preview_image_file(self, filepath):
    pixmap = QPixmap(filepath)
    if pixmap.isNull():
        self.image_view.setText("[Cannot load image]")
        self.preview_top.setCurrentWidget(self.image_view)
        return
    scaled = pixmap.scaled(
        max(1, self.preview_top.width()),
        max(1, self.preview_top.height()),
        Qt.KeepAspectRatio,
        Qt.SmoothTransformation
    )
    self.image_view.setPixmap(scaled)
    self.preview_top.setCurrentWidget(self.image_view)

def preview_text_file(self, filepath, max_chars=200000):
    try:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read(max_chars)
            if not content.strip():
                self.text_view.setText("[Empty file]")
            else:
                self.text_view.setText(content)
    except Exception as e:
        print("Text preview error:", e)
        self.text_view.setText("[Cannot read file]")
    self.preview_top.setCurrentWidget(self.text_view)

def preview_document_file(self, self_obj, filepath, max_lines=200):
    try:
        if filepath.lower().endswith(".docx"):
            from docx import Document
            doc = Document(filepath)
            text = []
            for paragraph in doc.paragraphs:
                text.append(paragraph.text)
                if len(text) >= max_lines:
                    break
            text = "\n".join(text)
            self.text_view.setText(text if text.strip() else "[Empty Word document]")
            self.preview_top.setCurrentWidget(self.text_view)
            return
        if filepath.lower().endswith(".xlsx"):
            wb = openpyxl.load_workbook(filepath, read_only=True)
            ws = wb.active
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                rows.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
                if i + 1 >= 50:
                    rows.append("... [Truncated]")
                    break
            self.text_view.setText("\n".join(rows) if rows else "[Empty Excel sheet]")
            self.preview_top.setCurrentWidget(self.text_view)
            return
    except Exception as e:
        print("Document preview error:", e)
        self.text_view.setText("[Cannot preview document]")
        self.preview_top.setCurrentWidget(self.text_view)
